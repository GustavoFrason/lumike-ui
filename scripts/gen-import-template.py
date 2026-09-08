"""
Gera os modelos .xlsx que o usuário baixa, preenche e sobe no admin:

  public/modelos/modelo-importacao-zarpellon.xlsx  -> /admin/compras/nova
  public/modelos/modelo-importacao-clientes.xlsx   -> /admin/clientes

Cada layout TEM que bater com o parser do backend correspondente
(`parseWorkbook` em lumike-api/src/{purchase-import,customer-import}/*.service.ts):

  - primeira aba, linha 1 = cabeçalho (o parser SEMPRE ignora a linha 1)
  - dados a partir da linha 2, lidos por POSIÇÃO de coluna
  - colunas depois da última útil são ignoradas (é onde ficam as instruções)

Rodar:  python scripts/gen-import-template.py   (precisa de openpyxl)
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MODELOS_DIR = Path(__file__).resolve().parent.parent / "public" / "modelos"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="C9A227")  # lumilee-gold
EXAMPLE_FONT = Font(italic=True, color="8A8A8A")


def _write_sheet(ws, headers, example, instructions, text_cols, widths, instr_col):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for col, value in enumerate(example, start=1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.font = EXAMPLE_FONT
        if col in text_cols:
            cell.number_format = "@"

    ws.cell(row=1, column=instr_col, value=instructions[0]).font = Font(bold=True)
    for i, line in enumerate(instructions[1:], start=2):
        ws.cell(row=i, column=instr_col, value=line)

    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def build_zarpellon() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Compra Zarpellon"
    _write_sheet(
        ws,
        headers=["Produto", "Descrição", "Qtd.", "Valor Base"],
        example=[
            "1171530601518",
            "ANEL PEQUENO ZIRCONIA REDONDO 3 GARRA RIVIERA 3MM BAN. DOURADO /INCOLOR /TAM. 18",
            10,
            8.85,
        ],
        instructions=[
            "COMO PREENCHER (apague esta coluna se quiser):",
            "1. Uma linha por produto, a partir da linha 2.",
            "2. Produto = código da peça na Zarpellon (coluna A).",
            "3. Descrição = nome/descrição da peça (coluna B).",
            "4. Qtd. = quantidade comprada. Valor Base = custo unitário (coluna D).",
            "5. A linha 2 é só um EXEMPLO — apague e coloque os seus itens.",
            "6. Não mexa na linha 1 (cabeçalho).",
        ],
        text_cols={1},  # coluna A (código) como texto: não perde zero à esquerda
        widths={"A": 18, "B": 70, "C": 8, "D": 12, "F": 62},
        instr_col=6,  # F
    )
    ws.cell(row=2, column=4).number_format = "0.00"
    out = MODELOS_DIR / "modelo-importacao-zarpellon.xlsx"
    wb.save(out)
    return out


def build_customers() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    _write_sheet(
        ws,
        headers=[
            "Nome",
            "Email",
            "Telefone",
            "CPF",
            "CEP",
            "Endereço",
            "Cidade",
            "Estado",
            "Observações",
        ],
        example=[
            "Maria Oliveira",
            "maria.oliveira@email.com",
            "(41) 99999-1234",
            "123.456.789-09",
            "80000-000",
            "Rua das Flores, 100",
            "Curitiba",
            "PR",
            "Cliente VIP",
        ],
        instructions=[
            "COMO PREENCHER (apague esta coluna se quiser):",
            "1. Uma linha por cliente, a partir da linha 2.",
            "2. Só o Nome (coluna A) é obrigatório — o resto é opcional.",
            "3. Telefone/CPF podem ir com ou sem máscara.",
            "4. Estado = a sigla da UF (ex: PR, SP).",
            "5. Cliente que já existe (mesmo e-mail/CPF/telefone) é ignorado.",
            "6. A linha 2 é só um EXEMPLO — apague e coloque os seus clientes.",
            "7. Não mexa na linha 1 (cabeçalho).",
        ],
        text_cols={3, 4, 5},  # Telefone, CPF, CEP como texto
        widths={
            "A": 24,
            "B": 28,
            "C": 18,
            "D": 18,
            "E": 12,
            "F": 30,
            "G": 18,
            "H": 8,
            "I": 24,
            "K": 60,
        },
        instr_col=11,  # K (o parser só lê A–I)
    )
    out = MODELOS_DIR / "modelo-importacao-clientes.xlsx"
    wb.save(out)
    return out


def main() -> None:
    MODELOS_DIR.mkdir(parents=True, exist_ok=True)
    for out in (build_zarpellon(), build_customers()):
        print(f"gerado: {out}")


if __name__ == "__main__":
    main()
